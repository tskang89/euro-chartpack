# -*- coding: utf-8 -*-
"""ifo 업황지수를 ifo 연구소에서 직접 받는다.

ifo 는 API 를 내지 않는다. 홈페이지는 자바스크립트 봇 검사(Client Challenge)가
걸려 있어 문서를 받아도 껍데기만 온다. 그런데 **엑셀 파일 자체는 그 검사를
지나지 않는다** — 직접 주소로 부르면 그대로 내려온다.

    https://www.ifo.de/sites/default/files/secure/timeseries/gsk-e-YYYYMM.xlsx

달마다 주소가 바뀌므로 이번 달부터 거꾸로 짚어 처음 잡히는 것을 쓴다. 발표는
매달 말이라 달이 바뀐 직후에는 지난달 파일이 최신이다.

들어 있는 것:
  'ifo Business Climate' 시트  독일 전체 업황지수 (2015=100, 계절조정)
  'Sectors' 시트               제조업·서비스업 등 부문별

부문별은 **지수가 아니라 잔액(balance)** 이다. 좋다고 답한 비율에서 나쁘다고
답한 비율을 뺀 값이라 −30~+30 언저리를 오간다. 전체 지수(85~100)와 한 축에
놓으면 읽히지 않으므로 화면에서 따로 그린다. 단위가 다른 것을 같은 그림에
얹지 않는다.
"""

from __future__ import annotations

import datetime
import io
import time

import requests

URL = ("https://www.ifo.de/sites/default/files/secure/timeseries/"
       "gsk-e-{ym}.xlsx")
TIMEOUT = 90
LOOK_BACK = 4          # 이번 달부터 몇 달까지 거슬러 찾아볼 것인가
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")

# 'Sectors' 시트에서 업황지수(Business Climate)가 있는 열. 1부터 센다.
SECTOR_COL = {"man": 8, "srv": 11}      # 제조업, 서비스업


class IfoError(RuntimeError):
    pass


def _fetch(ym: str) -> bytes | None:
    """그 달 파일을 받는다. 없으면 None.

    없는 달을 부르면 봇 검사 페이지(HTML)가 200 으로 온다. 그래서 상태
    코드가 아니라 내용을 봐야 한다 — xlsx 는 zip 이라 'PK' 로 시작한다.
    """
    try:
        resp = requests.get(URL.format(ym=ym), timeout=TIMEOUT,
                            headers={"User-Agent": UA})
    except requests.RequestException:
        return None
    if resp.status_code != 200 or not resp.content.startswith(b"PK"):
        return None
    return resp.content


def workbook(today: datetime.date | None = None, log=print):
    """가장 최근 판을 받아 (엑셀, 기준연월) 로 돌려준다."""
    today = today or datetime.date.today()
    y, m = today.year, today.month
    for _ in range(LOOK_BACK):
        ym = f"{y}{m:02d}"
        raw = _fetch(ym)
        if raw:
            import openpyxl
            log(f"    ifo 원본 gsk-e-{ym}.xlsx ({len(raw):,}바이트)")
            return openpyxl.load_workbook(io.BytesIO(raw), data_only=True), ym
        time.sleep(0.5)
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    raise IfoError(f"최근 {LOOK_BACK}달치 파일을 하나도 받지 못했다")


def _month_key(cell) -> str | None:
    """' 09/2026' 또는 날짜 값을 'YYYY-MM' 으로."""
    if cell is None:
        return None
    if isinstance(cell, datetime.datetime):
        return f"{cell.year}-{cell.month:02d}"
    text = str(cell).strip()
    if "/" in text:
        mm, _, yy = text.partition("/")
        try:
            return f"{int(yy)}-{int(mm):02d}"
        except ValueError:
            return None
    return None


def _column(ws, col: int, first_row: int = 9) -> dict[str, float]:
    out: dict[str, float] = {}
    for row in ws.iter_rows(min_row=first_row, max_row=ws.max_row,
                            max_col=col, values_only=True):
        key = _month_key(row[0])
        if key is None:
            continue
        value = row[col - 1]
        if isinstance(value, (int, float)):
            out[key] = float(value)
    return out


def series(today: datetime.date | None = None,
           log=print) -> dict[str, dict[str, float]]:
    """{'all': {월: 지수}, 'man': {월: 잔액}, 'srv': {월: 잔액}}."""
    wb, _ym = workbook(today, log=log)
    out = {"all": _column(wb["ifo Business Climate"], 2)}
    sectors = wb["Sectors"]
    for key, col in SECTOR_COL.items():
        out[key] = _column(sectors, col)
    return out
